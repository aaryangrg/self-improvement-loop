#!/usr/bin/env python3
"""Parse task document files to text for the recipe-local read tool."""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
import pdfplumber


def parse_docx(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text:
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
        parts.append("")
    return "\n".join(parts)


def parse_pdf(path: Path) -> str:
    parts: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text)
            for table in page.extract_tables():
                for row in table:
                    parts.append("\t".join(cell or "" for cell in row))
                parts.append("")
    return "\n".join(parts)


def parse_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(path)
    parts: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts.append(f"=== Slide {slide_index} ===")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
        parts.append("")
    return "\n".join(parts)


def parse_xlsx(path: Path) -> str:
    sheets = pd.read_excel(path, sheet_name=None)
    parts: list[str] = []
    for name, frame in sheets.items():
        parts.append(f"=== Sheet: {name} ===")
        parts.append(frame.to_string(index=False))
    return "\n".join(parts)


PARSERS = {
    "docx": parse_docx,
    "pdf": parse_pdf,
    "pptx": parse_pptx,
    "xlsx": parse_xlsx,
}


def main() -> int:
    if len(sys.argv) != 3 or sys.argv[1] not in PARSERS:
        print(f"usage: {sys.argv[0]} {{{'|'.join(PARSERS)}}} <path>", file=sys.stderr)
        return 2

    file_type = sys.argv[1]
    path = Path(sys.argv[2])
    try:
        sys.stdout.write(PARSERS[file_type](path))
        return 0
    except Exception as error:
        print(f"{type(error).__name__}: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
